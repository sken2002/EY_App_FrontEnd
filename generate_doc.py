import os
try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    print("python-docx not installed.")
    exit(1)

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

doc.add_heading('Project Spider: System Architecture Overview', 0)

doc.add_paragraph('This document outlines the core technical structure of the Project Spider risk management platform, detailing both the data processing backend and the interactive frontend dashboard.')

doc.add_heading('1. Data Processing Pipeline (Backend)', level=1)
doc.add_paragraph('Located in the "pipeline/" directory, this subsystem is responsible for transforming raw operational data into structured risk metrics.')

doc.add_heading('pipeline/compute_risks.py', level=2)
doc.add_paragraph('Functions as the primary calculation engine. It ingests flat data files (e.g., project budgets, milestones, and supplier metrics) and computes severity scores across six distinct risk dimensions (Cost, Schedule, Operational, Supplier, Cashflow, and Data Quality).')

doc.add_heading('pipeline/export_state.py', level=2)
doc.add_paragraph('Constructs the dependency graph. It maps relationships between Work Packages, contracts, and milestones, eventually packaging the computed risk metrics and node connections into a single comprehensive JSON state file.')

doc.add_heading('public/data/state.json', level=2)
doc.add_paragraph('Serves as the static database for the application. This file contains the complete, pre-calculated graph topology and risk scores used by the frontend.')

doc.add_heading('2. Application Interface (Frontend)', level=1)
doc.add_paragraph('Located in the "frontend/" directory, this subsystem is a Next.js web application providing an interactive dashboard and real-time simulation capabilities.')

doc.add_heading('src/app/page.tsx', level=2)
doc.add_paragraph('The root routing and layout controller. It manages the global application state, including contextual variables (Project Archetype, Time Horizon) and orchestrates navigation between the Macro, Portfolio, and Entity-level drill-down views.')

doc.add_heading('src/components/layout/LeftPanel.tsx', level=2)
doc.add_paragraph('The executive summary component. It aggregates portfolio-wide risk scores, providing high-level visibility into the health of the entire project across the six core dimensions.')

doc.add_heading('src/components/canvas/CenterCanvas.tsx & MacroHeatmap.tsx', level=2)
doc.add_paragraph('The primary visualization components. CenterCanvas renders the interactive node-based dependency graph, allowing users to trace relationships between tasks. The MacroHeatmap provides a clustered overview of project workstreams, highlighting structural weaknesses before drilling down.')

doc.add_heading('src/components/layout/RightPanel.tsx', level=2)
doc.add_paragraph('The analysis and simulation hub. It renders dynamic Radar charts for specific project entities and provides interactive controls that allow users to simulate metric changes ("what-if" scenarios).')

doc.add_heading('src/lib/riskEngine/', level=2)
doc.add_paragraph('The client-side mathematical engine. When users simulate scenarios in the Right Panel, the algorithms in this directory dynamically recalculate the Composite Risk Index and traverse the dependency graph to compute the downstream financial exposure (the "Blast Radius").')

doc.add_heading('src/app/api/narrative/route.ts', level=2)
doc.add_paragraph('The AI integration layer. This endpoint bridges the deterministic simulation engine with a large language model. It securely passes the calculated risk state and global context variables to the AI, returning synthesized, context-aware mitigation strategies directly into the dashboard.')

output_path = r'c:\Users\SKEN\Desktop\LBS\EY\Data\project-spider\Project_Spider_Architecture.docx'
doc.save(output_path)
print(f"Document saved to {output_path}")
